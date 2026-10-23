"""
Webinar Dev Bot — PPTX builder (Canva delivery, Route A).

Turns a slides.json manifest into a native 16:9 PowerPoint file with REAL,
editable text boxes and the slide images placed in their slots. That PPTX is
what gets dragged into Canva (works on every Canva plan, nothing goes public).

Why not export the HTML deck to PDF? A PDF of the HTML lands in Canva as flat
pictures. Real text boxes survive the import and stay editable.

Usage:
    python3 build_pptx.py path/to/slides.json [--out deck.pptx] [--strict]

Exit codes:
    0  built, no warnings
    1  bad manifest or cannot write the file (plain-language message printed)
    2  built, but with warnings (only when --strict is passed)

Requires: pip install python-pptx Pillow
Manifest format: references/canva-delivery.md ("slides.json manifest").
"""

import argparse
import io
import json
import math
import os
import re
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- Geometry: exact 16:9 (PowerPoint/Canva "Widescreen") -------------------
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
BAR_H = Inches(0.5)            # bottom bar, like the HTML deck's 48px bar
CONTENT_H = SLIDE_H - BAR_H
PAD = Inches(0.6)
MAX_IMAGE_EDGE = 2400          # px; keeps decks small enough for Canva import

# Layout -> (image column fraction, image side). Mirrors references/design-system.md.
LAYOUTS = {
    "split":         (0.42, "left"),
    "dark":          (0.42, "left"),
    "split-reverse": (0.45, "right"),
    "center":        (0.0, None),
    "stack":         (0.0, None),
}

DEFAULT_BRAND = {
    "primary": "#4A90D9",
    "primary_light": "#6BB3F0",
    "secondary": "#C0392B",
    "accent": "#C9A84C",
    "text": "#222222",
    "heading_text": "#000000",
    "muted": "#666666",
    "background": "#FFFFFF",
    "dark_background": "#000000",
    "heading_font": "Playfair Display",
    "body_font": "Source Sans 3",
}

# Google Fonts name -> the name Canva/Office actually list.
FONT_ALIASES = {"Source Sans 3": "Source Sans Pro"}

HEX_RE = re.compile(r"^#?[0-9A-Fa-f]{6}$")
LAYOUT_NAMES = ", ".join(LAYOUTS)


class ManifestError(Exception):
    """Raised for input problems; main() prints the message and exits 1."""


# ---------------------------------------------------------------------------
# Small helpers

def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def font_name(name):
    return FONT_ALIASES.get(name, name)


def safe_path(base, rel, warnings, label):
    """Resolve a manifest path and refuse anything outside the project folder.
    Returns a Path, or None (caller shows a placeholder)."""
    if not rel:
        return None
    p = (base / str(rel)).resolve()
    if not p.is_relative_to(base.resolve()):
        warnings.append(f"{label}: '{rel}' is outside the project folder, skipped")
        return None
    if not p.exists():
        warnings.append(f"{label}: file not found: {rel}")
        return None
    return p


def fit_size(text, box_w, box_h, start, floor, serif=False, line_gap=1.2, para_gap=0):
    """Pick the largest font size (pt) that fits the box, by estimating glyph width.
    Canva ignores PowerPoint's shrink-on-overflow, so we size text ourselves.
    Returns (size, fits)."""
    width_pt = box_w / 12700          # EMU -> points
    height_pt = box_h / 12700
    em = 0.54 if serif else 0.5
    paragraphs = [p for p in str(text).split("\n")] or [""]
    size = start
    while True:
        chars_per_line = max(1, int(width_pt / (size * em)))
        lines = sum(max(1, math.ceil(len(p) / chars_per_line)) for p in paragraphs)
        needed = lines * size * line_gap + para_gap * len(paragraphs)
        if needed <= height_pt:
            return size, True
        if size <= floor:
            return floor, False
        size -= 2


def text_height(text, box_w, size, serif=False, line_gap=1.2, para_gap=0):
    width_pt = box_w / 12700
    em = 0.54 if serif else 0.5
    chars_per_line = max(1, int(width_pt / (size * em)))
    paragraphs = str(text).split("\n")
    lines = sum(max(1, math.ceil(len(p) / chars_per_line)) for p in paragraphs)
    return Pt(lines * size * line_gap + para_gap * len(paragraphs) + 6)


def add_runs(paragraph, text, brand, size, font, color, bold=False, italic=False, hl_color=None):
    """**double-asterisk** words become bold + primary color (the deck's .highlight).
    Pass hl_color on colored backgrounds (cards), where primary would vanish."""
    for part in re.split(r"(\*\*[^*]+\*\*)", str(text)):
        if not part:
            continue
        hl = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run()
        run.text = part[2:-2] if hl else part
        run.font.size = Pt(size)
        run.font.name = font_name(font)
        run.font.bold = bold or hl
        run.font.italic = italic
        run.font.color.rgb = rgb((hl_color or brand["primary"]) if hl else color)


def new_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return tf


def solid_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


# ---------------------------------------------------------------------------
# Images

def prepared_image(path, box_w, box_h):
    """Return a BytesIO of the image cropped like CSS object-fit:cover;
    object-position:top, shrunk to MAX_IMAGE_EDGE. Pre-cropping (instead of PPTX
    crop values) keeps Canva and Google Slides from mis-rendering it."""
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        box_ratio = box_w / box_h
        iw, ih = im.size
        if iw / ih > box_ratio:                 # too wide -> trim sides evenly
            new_w = int(ih * box_ratio)
            left = (iw - new_w) // 2
            im = im.crop((left, 0, left + new_w, ih))
        else:                                   # too tall -> keep the top
            new_h = int(iw / box_ratio)
            im = im.crop((0, 0, iw, new_h))
        im.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
        buf = io.BytesIO()
        if im.mode in ("RGBA", "LA", "P"):
            im.save(buf, "PNG")
        else:
            im.convert("RGB").save(buf, "JPEG", quality=88)
        buf.seek(0)
        return buf


def place_image(slide, path, x, y, w, h, warnings, label):
    try:
        slide.shapes.add_picture(prepared_image(path, w, h), x, y, w, h)
        return True
    except Exception as exc:  # unreadable/corrupt image -> placeholder, keep building
        warnings.append(f"{label}: could not read image ({exc.__class__.__name__}), used a placeholder")
        return False


def load_logo(path, warnings):
    """Read + normalize the logo once. Unreadable logo -> warning, no logo (never a crash)."""
    if not path:
        return None
    try:
        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im)
            im.thumbnail((800, 800))
            buf = io.BytesIO()
            im.save(buf, "PNG")
            return buf.getvalue()
    except Exception as exc:
        warnings.append(f"logo: could not read it as an image ({exc.__class__.__name__}), built without a logo")
        return None


def placeholder(slide, x, y, w, h, label, brand):
    shape = solid_rect(slide, x, y, w, h, brand["primary"])
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, f"[YOUR IMAGE: {label}]", brand, 16, brand["body_font"], "#FFFFFF")


# ---------------------------------------------------------------------------
# Content blocks

def render_content(slide, s, brand, x, y, w, h, dark, centered, warnings, label):
    """Title, subtitle, then one of: cards, table, body/bullets."""
    heading_color = "#FFFFFF" if dark else brand["heading_text"]
    body_color = "#EEEEEE" if dark else brand["text"]
    muted = "#CCCCCC" if dark else brand["muted"]
    align = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT

    title = str(s.get("title", ""))
    subtitle = str(s.get("subtitle", ""))
    body = str(s.get("body", ""))
    bullets = [str(b) for b in s.get("bullets", [])]
    body_text = "\n".join(([body] if body else []) + [f"•  {b}" for b in bullets])

    # Size the title first (max 2.2in tall)
    title_size, ok = fit_size(title, w, Inches(2.2), s.get("title_size", 40), 26, serif=True)
    if title and not ok:
        warnings.append(f"{label}: title is too long to fit, shorten it")
    title_h = text_height(title, w, title_size, serif=True) if title else 0
    sub_h = text_height(subtitle, w, 22, serif=True) if subtitle else 0
    line_block = Inches(0.35) if title else 0

    # For center layout, measure the whole block so it sits in the middle like the HTML.
    remaining_for_body = h - title_h - line_block - sub_h
    body_size = s.get("body_size", 22)
    if body_text and not s.get("cards") and not s.get("table"):
        body_size, ok = fit_size(body_text, w, remaining_for_body, body_size, 14, para_gap=10)
        if not ok:
            warnings.append(f"{label}: body text overflows the slide, trim it or split the slide")
    body_h = text_height(body_text, w, body_size, para_gap=10) if body_text else 0

    cursor = y
    if centered and not s.get("cards") and not s.get("table"):
        block = title_h + line_block + sub_h + body_h
        cursor = y + max(0, (h - block) // 2)

    if title:
        # Centered slides: bottom-anchor so the title sits snug on its accent line
        # (any misjudged wrap spills into the empty space above). Elsewhere top-anchor,
        # so a longer-than-estimated title can never climb past the slide's top edge.
        anchor = MSO_ANCHOR.BOTTOM if centered else MSO_ANCHOR.TOP
        tf = new_textbox(slide, x, cursor, w, title_h, anchor=anchor)
        p = tf.paragraphs[0]
        p.alignment = align
        add_runs(p, title, brand, title_size, brand["heading_font"], heading_color, bold=True)
        cursor += title_h
        line_w = Inches(1.25)
        line_x = x + (w - line_w) // 2 if centered else x + Inches(0.08)
        solid_rect(slide, line_x, cursor + Inches(0.1), line_w, Inches(0.055), brand["primary"],
                   MSO_SHAPE.ROUNDED_RECTANGLE)
        cursor += line_block

    if subtitle:
        tf = new_textbox(slide, x, cursor, w, sub_h)
        p = tf.paragraphs[0]
        p.alignment = align
        add_runs(p, subtitle, brand, 22, brand["heading_font"], muted, italic=True)
        cursor += sub_h

    remaining = y + h - cursor
    if s.get("cards"):
        render_cards(slide, s["cards"], brand, x, cursor, w, remaining, warnings, label)
    elif s.get("table"):
        render_table(slide, s, brand, x, cursor, w, remaining, warnings, label)
    elif body_text:
        tf = new_textbox(slide, x, cursor, w, max(body_h, remaining))
        first = True
        for line in [ln for ln in body.split("\n") if ln.strip()]:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            add_runs(p, line, brand, body_size, brand["body_font"], body_color)
            p.space_after = Pt(10)
        for item in bullets:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            add_runs(p, f"•  {item}", brand, body_size, brand["body_font"], body_color)
            p.space_after = Pt(6)


def render_cards(slide, cards, brand, x, y, w, h, warnings, label):
    """Belief / If-All-I-Did-Was / Two Paths boxes. 4 cards = 2x2, otherwise rows of 3."""
    n = len(cards)
    per_row = 2 if n == 4 else min(3, n)
    rows = math.ceil(n / per_row)
    gap = Inches(0.25)
    top = Inches(0.2)
    avail = max(h - top, Inches(0.9))
    if h - top < Inches(0.9):
        warnings.append(f"{label}: the title leaves no room for the cards, shorten the title")
    card_w = (w - gap * (per_row - 1)) // per_row
    # One row of cards may use the full height (long lists); stacked rows stay compact.
    card_h = min(Inches(4.2) if rows == 1 else Inches(2.4), (avail - gap * (rows - 1)) // rows)
    if card_h < Inches(0.9):
        warnings.append(f"{label}: {n} cards don't fit well, consider splitting the slide")
    palette = [brand["primary"], brand["secondary"]]   # gold/accent fails white-text contrast
    for i, card in enumerate(cards):
        if not isinstance(card, dict):
            card = {"title": str(card)}
        r, c = divmod(i, per_row)
        cx = x + c * (card_w + gap)
        cy = y + top + r * (card_h + gap)
        color = card.get("color") or palette[i % len(palette)]
        shape = solid_rect(slide, cx, cy, card_w, card_h, color, MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(0.15)
        ctitle, cbody = str(card.get("title", "")), str(card.get("body", ""))
        tsize, _ = fit_size(ctitle, card_w - Inches(0.3), card_h * 0.3, 22, 14, serif=True)
        bsize, ok = fit_size(cbody, card_w - Inches(0.3), card_h * 0.65, 18, 11)
        if cbody and not ok:
            warnings.append(f"{label}: card {i + 1} text is too long")
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, ctitle, brand, tsize, brand["heading_font"], "#FFFFFF", bold=True, hl_color="#FFFFFF")
        if cbody:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            add_runs(p2, cbody, brand, bsize, brand["body_font"], "#FFFFFF", hl_color="#FFFFFF")


def render_table(slide, s, brand, x, y, w, h, warnings, label):
    """Value-stack / timeline tables. Two columns default to 70/30 with the value
    right-aligned in primary (the deck's .stack-table). Set "table_header": true
    to style row 1 as a header, "table_total": true to bold the last row."""
    rows = [[str(c) for c in row] for row in s["table"]]
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    widths = list(s.get("col_widths") or ([0.7, 0.3] if n_cols == 2 else [1 / n_cols] * n_cols))[:n_cols]
    if len(widths) < n_cols:      # fill missing columns with an equal share of what's left
        rest = max(0.0, 1 - sum(widths))
        widths += [rest / (n_cols - len(widths)) or 1 / n_cols] * (n_cols - len(widths))
    total_w = sum(widths) or 1
    widths = [x / total_w for x in widths]          # always fits the box
    row_h = Inches(0.5)
    if row_h * n_rows > h:
        warnings.append(f"{label}: table has more rows than fit, split it across slides")
        row_h = max(Inches(0.32), h // n_rows)
    header, total = s.get("table_header", False), s.get("table_total", False)
    table = slide.shapes.add_table(n_rows, n_cols, x, y + Inches(0.2), w, row_h * n_rows).table
    for i, frac in enumerate(widths[:n_cols]):
        table.columns[i].width = Emu(int(w * frac))
    for r, row in enumerate(rows):
        is_header = header and r == 0
        is_total = total and r == n_rows - 1
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            text = row[c] if c < len(row) else ""
            value_col = n_cols == 2 and c == 1
            p.alignment = PP_ALIGN.RIGHT if value_col else PP_ALIGN.LEFT
            if is_header:
                color, fill = "#FFFFFF", brand["primary"]
            else:
                color = brand["primary"] if value_col else brand["text"]
                fill = "#F8F9FA" if r % 2 else "#FFFFFF"
            size = 20 if is_total else 16
            add_runs(p, text, brand, size, brand["body_font"], color,
                     bold=is_header or is_total or value_col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)


def render_bottom_bar(slide, n, total, brand, logo):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SLIDE_H - BAR_H, SLIDE_W, BAR_H)
    bar.line.fill.background()
    bar.shadow.inherit = False
    bar.fill.gradient()
    bar.fill.gradient_angle = 0
    stops = bar.fill.gradient_stops
    stops[0].color.rgb = rgb(brand["primary"])
    stops[1].color.rgb = rgb(brand["primary_light"])
    if logo:
        pic = slide.shapes.add_picture(io.BytesIO(logo), Inches(0.3), SLIDE_H - BAR_H + Inches(0.08),
                                       height=BAR_H - Inches(0.16))
        if pic.width > Inches(2.0):          # cap very wide logos
            ratio = Inches(2.0) / pic.width
            pic.width, pic.height = Inches(2.0), int(pic.height * ratio)
    tf = new_textbox(slide, SLIDE_W - Inches(1.8), SLIDE_H - BAR_H, Inches(1.5), BAR_H,
                     anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_runs(p, f"{n} / {total}", brand, 12, brand["body_font"], "#FFFFFF", bold=True)


# ---------------------------------------------------------------------------
# Manifest validation + build

def validate(data):
    if not isinstance(data, dict):
        raise ManifestError("slides.json must be a JSON object with a \"slides\" list.")
    slides = data.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ManifestError("slides.json needs a non-empty \"slides\" list.")
    raw_brand = data.get("brand") or {}
    if not isinstance(raw_brand, dict):
        raise ManifestError('"brand" must be an object like {"primary": "#4A90D9"}.')
    brand = {**DEFAULT_BRAND, **raw_brand}
    if "primary_light" not in raw_brand and isinstance(brand.get("primary"), str) and HEX_RE.match(brand["primary"]):
        # bottom-bar gradient end: the brand primary mixed ~35% toward white (never the default blue)
        h = brand["primary"].lstrip("#")
        rgb_ = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
        brand["primary_light"] = "#" + "".join(f"{round(c + (255 - c) * 0.35):02X}" for c in rgb_)
    for key, val in brand.items():
        if key.endswith("_font"):
            if not isinstance(val, str) or not val.strip():
                raise ManifestError(f"brand.{key} must be a font name (got {val!r}).")
            continue
        if not isinstance(val, str) or not HEX_RE.match(val):
            raise ManifestError(f"brand.{key} must be a hex color like #4A90D9 (got {val!r}).")
    if data.get("logo") is not None and not isinstance(data["logo"], str):
        raise ManifestError('"logo" must be a file path.')
    for i, s in enumerate(slides, start=1):
        if not isinstance(s, dict):
            raise ManifestError(f"slide {i} must be an object.")
        for key in ("title_size", "body_size"):
            if key in s:
                v = s[key]
                if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or not 8 <= v <= 120:
                    raise ManifestError(f"slide {i}: {key} must be a number from 8 to 120 (got {v!r}).")
        for key in ("layout", "image", "image_description"):
            if s.get(key) is not None and not isinstance(s[key], str):
                raise ManifestError(f"slide {i}: {key} must be text.")
        if "bullets" in s and not isinstance(s["bullets"], list):
            raise ManifestError(f"slide {i}: bullets must be a list.")
        if "cards" in s:
            if not isinstance(s["cards"], list) or not s["cards"]:
                raise ManifestError(f"slide {i}: cards must be a non-empty list.")
            for j, c in enumerate(s["cards"], start=1):
                col = c.get("color") if isinstance(c, dict) else None
                if col is not None and (not isinstance(col, str) or not HEX_RE.match(col)):
                    raise ManifestError(f"slide {i}, card {j}: color must be a hex color like #4A90D9 (got {col!r}).")
        if "table" in s:
            t = s["table"]
            if not isinstance(t, list) or not t or not all(isinstance(r, list) and r for r in t):
                raise ManifestError(f"slide {i}: table must be a list of non-empty rows.")
        if "col_widths" in s:
            cw = s["col_widths"]
            if (not isinstance(cw, list) or not cw
                    or not all(isinstance(x, (int, float)) and not isinstance(x, bool)
                               and math.isfinite(x) and x > 0 for x in cw)):
                raise ManifestError(f"slide {i}: col_widths must be a list of positive numbers like [0.7, 0.3].")
    return brand, slides


def build(manifest_path, out_path):
    manifest_path = Path(manifest_path)
    base = manifest_path.parent.resolve()
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise ManifestError(f"Can't find {manifest_path}.")
    except json.JSONDecodeError as exc:
        raise ManifestError(f"slides.json isn't valid JSON (line {exc.lineno}): {exc.msg}.")
    brand, slides = validate(data)
    warnings = []
    logo = load_logo(safe_path(base, data.get("logo"), warnings, "logo"), warnings)

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    for i, s in enumerate(slides, start=1):
        label = f"slide {i}"
        layout = s.get("layout", "split")
        if layout not in LAYOUTS:
            warnings.append(f"{label}: unknown layout '{layout}' (use one of: {LAYOUT_NAMES}), used 'center'")
            layout = "center"
        img_frac, side = LAYOUTS[layout]
        dark = layout == "dark" or s.get("dark") is True

        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(brand["dark_background"] if dark else brand["background"])

        if img_frac:
            img_w = Emu(int(SLIDE_W * img_frac))
            img_x = 0 if side == "left" else SLIDE_W - img_w
            img = safe_path(base, s.get("image"), warnings, label)
            if not (img and place_image(slide, img, img_x, 0, img_w, CONTENT_H, warnings, label)):
                placeholder(slide, img_x, 0, img_w, CONTENT_H, s.get("image_description", "image"), brand)
            text_x = img_w + PAD if side == "left" else PAD
            text_w = SLIDE_W - img_w - PAD * 2
            render_content(slide, s, brand, text_x, PAD, text_w, CONTENT_H - PAD * 2,
                           dark, False, warnings, label)
        else:
            centered = layout == "center"
            text_w = min(SLIDE_W - PAD * 2, Inches(11)) if centered else SLIDE_W - Inches(1.6)
            text_x = (SLIDE_W - text_w) // 2
            render_content(slide, s, brand, text_x, PAD, text_w, CONTENT_H - PAD * 2,
                           dark, centered, warnings, label)

        render_bottom_bar(slide, i, len(slides), brand, logo)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])

    # Write to a temp file first so a crash never leaves a half-written deck.
    out_path = Path(out_path)
    if not out_path.parent.is_dir():
        raise ManifestError(f"The folder for {out_path} doesn't exist.")
    tmp = None
    try:
        fd, tmp = tempfile.mkstemp(suffix=".pptx", dir=out_path.parent)
        os.close(fd)
        prs.save(tmp)
        os.replace(tmp, out_path)
    except PermissionError:
        raise ManifestError(f"Can't write {out_path.name}. If it's open in PowerPoint or Keynote, close it and run again.")
    finally:
        if tmp and Path(tmp).exists():
            Path(tmp).unlink(missing_ok=True)
    return len(slides), warnings


def main():
    ap = argparse.ArgumentParser(description="Build a Canva-ready PPTX from slides.json")
    ap.add_argument("manifest")
    ap.add_argument("--out", help="Output .pptx path (default: deck.pptx next to the manifest)")
    ap.add_argument("--strict", action="store_true", help="Exit 2 if there are any warnings")
    args = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    out = Path(args.out) if args.out else Path(args.manifest).with_name("deck.pptx")
    try:
        count, warnings = build(args.manifest, out)
    except ManifestError as exc:
        print(f"ERROR: {exc}")
        sys.exit(1)
    except Exception as exc:   # never show a raw traceback to a non-developer
        print(f"ERROR: something unexpected went wrong while building ({exc.__class__.__name__}: {str(exc)[:200]}).")
        print("Run it again; if it repeats, share this message with whoever supports your kit.")
        sys.exit(1)
    for w in warnings:
        print(f"WARNING: {w}")
    print(f"Built {count} slides -> {out} ({len(warnings)} warning{'s' if len(warnings) != 1 else ''})")
    sys.exit(2 if (args.strict and warnings) else 0)


if __name__ == "__main__":
    main()
